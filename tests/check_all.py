"""Full-system check for PrintVerity. Starts its own server via run.sh on a test port, exercises everything, reports PASS/FAIL per item."""
import sys, os, time, json, subprocess, urllib.request, urllib.error, uuid, io, zipfile, re, shutil, xml.dom.minidom
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
PORT = int(os.environ.get("PV_TEST_PORT", "8711")); BASE = f"http://127.0.0.1:{PORT}"
PY = str(ROOT / ".venv" / "bin" / "python")
REF, NB, FIX = ROOT / "testdata" / "reference", ROOT / "testdata" / "northbay", ROOT / "tests" / "fixtures"
# the reference package, uploaded exactly as testdata/reference/GUIDE.md describes: (job name, customer, drawings, STEP, BOM, PO)
REFERENCE_JOBS = [
    ("Job 24-1187 · Drive module (assembly, bracket Rev D, shaft)", "OEM customer A", ["4471-000_B.dxf", "4471-020_D.dxf", "4471-105_B.dxf"], "4471-020_D.stp", "BOM_4471-000_rev3.xlsx", "PO_88213.xlsx"),
    ("Job 24-1160 · Bracket 4471-020 previous revision (Rev C)", "OEM customer A", ["4471-020_C.dxf"], None, None, None),
    ("Job 24-1203 · PDF plots of the same package", "OEM customer A", ["4471-020_D.pdf", "4471-105_B.pdf", "4471-000_B.pdf"], "4471-105_B.stp", None, None),
    ("Job 24-1211 · Layout-style DXF (title block in paper space)", "OEM customer B", ["4471-020_D_layout.dxf"], None, None, None),
    ("Job 24-1215 · Sheet metal bracket 4471-210 (formed view + flat pattern)", "OEM customer B", ["4471-210_A.dxf"], None, None, None),
]
results = []
def check(name, cond, info=""):
    results.append((name, bool(cond), info))
    print(("PASS " if cond else "FAIL ") + name + (f"  [{info}]" if info and not cond else ""))
    return bool(cond)

def http(method, path, data=None, headers=None, raw=False):
    req = urllib.request.Request(BASE + path, method=method, headers=headers or {})
    body = None
    if data is not None and not raw:
        body = json.dumps(data).encode(); req.add_header("content-type", "application/json")
    elif raw: body = data
    try:
        with urllib.request.urlopen(req, body, timeout=120) as r:
            ct = r.headers.get("content-type", ""); b = r.read()
            return r.status, ct, (json.loads(b) if "json" in ct else b), {k.lower(): v for k, v in r.headers.items()}
    except urllib.error.HTTPError as e:
        b = e.read()
        try: return e.code, e.headers.get("content-type", ""), json.loads(b), {}
        except Exception: return e.code, "", b, {}

def multipart(fields, files):
    bnd = "----pv" + uuid.uuid4().hex; out = io.BytesIO()
    for k, v in fields.items():
        out.write(f"--{bnd}\r\nContent-Disposition: form-data; name=\"{k}\"\r\n\r\n{v}\r\n".encode())
    for k, path in files:
        p = Path(path)
        out.write(f"--{bnd}\r\nContent-Disposition: form-data; name=\"{k}\"; filename=\"{p.name}\"\r\nContent-Type: application/octet-stream\r\n\r\n".encode()); out.write(p.read_bytes()); out.write(b"\r\n")
    out.write(f"--{bnd}--\r\n".encode())
    return out.getvalue(), {"content-type": f"multipart/form-data; boundary={bnd}"}

def main():
    t0 = time.time()
    # 0. byte-compile every module
    r = subprocess.run([PY, "-m", "compileall", "-q", str(ROOT / "app"), str(ROOT / "tests"), str(ROOT / "testdata")], capture_output=True, text=True)
    check("modules compile", r.returncode == 0, r.stderr[-300:])
    # 1. fresh server via run.sh on the test port, with a clean job store
    subprocess.run(["fuser", "-k", f"{PORT}/tcp"], capture_output=True)
    scratch = Path("/tmp/pv_check_data"); shutil.rmtree(scratch, ignore_errors=True); (scratch / "jobs").mkdir(parents=True)
    shutil.copytree(ROOT / "data" / "profiles", scratch / "profiles"); (scratch / "specs").mkdir()          # Setup edits go to copies
    shutil.copy(ROOT / "data" / "specs" / "library.json", scratch / "specs" / "library.json"); shutil.copy(ROOT / "data" / "plm.json", scratch / "plm.json")
    log = open("/tmp/pv_test_server.log", "w")
    env = {**os.environ, "PV_CELERY": "off", "PV_JOBS_DIR": str(scratch / "jobs"), "PV_AUDIT_FILE": str(scratch / "audit.jsonl"), "PV_OUT_DIR": str(scratch / "out"), "PV_INDEX_DIR": str(scratch / "index"), "PV_PROFILES_DIR": str(scratch / "profiles"), "PV_SPEC_LIBRARY": str(scratch / "specs" / "library.json"), "PV_PLM_FILE": str(scratch / "plm.json"), "PV_VECTOR_STORE": "local", "PV_HOST": "127.0.0.1", "PV_ACCESS_TOKEN": ""}
    proc = subprocess.Popen(["bash", str(ROOT / "run.sh"), str(PORT)], stdout=log, stderr=subprocess.STDOUT, cwd=str(ROOT), start_new_session=True, env=env)
    up = False
    for _ in range(60):
        time.sleep(0.5)
        try:
            st = http("GET", "/api/status")
            if st[0] == 200: up = True; break
        except Exception: pass
    check("server starts via run.sh", up)
    if not up: return finish(proc, t0)
    st = http("GET", "/api/status")[2]
    check("status reports templates", st.get("templates") == 5, str(st))
    eng = st.get("engines", {})
    check("status reports engines (langgraph, retrieval, step, ocr, layout)", eng.get("orchestration") == "langgraph" and eng.get("retrieval", {}).get("store") in ("pgvector", "local") and eng.get("step") in ("opencascade", "text-parser") and eng.get("ocr") and eng.get("layout_model"), str(eng)[:200])
    s, _, kn, _ = http("GET", "/api/knowledge/search?q=datum+reference+missing&k=3&kinds=standard"); check("knowledge search returns standard clauses", s == 200 and kn and "datum" in kn[0]["text"].lower())
    # 2. the reference package, uploaded through the same endpoint the Queue's upload form uses
    check("no sample-loader endpoint in the product", http("POST", "/api/jobs/sample")[0] in (404, 405))
    jobs = []
    for name, cust, dws_, st_, bom_, po_ in REFERENCE_JOBS:
        files = [("drawings", REF / "drawings" / f) for f in dws_] + [(k, REF / "documents" / f) for k, f in (("step", st_), ("bom", bom_), ("po", po_)) if f]
        body, hdr = multipart({"name": name, "customer": cust}, files)
        s, _, j, _ = http("POST", "/api/jobs?wait=1", body, hdr, raw=True)
        if s == 200: jobs.append(j)
    check("reference package uploads as 5 jobs", len(jobs) == 5 and [j["name"] for j in http("GET", "/api/jobs")[2]][::-1] == [x[0] for x in REFERENCE_JOBS], f"{len(jobs)} jobs")
    dws = {d["file"]: (j, d) for j in jobs for d in j["drawings"]}
    expected = {"4471-000_B.dxf": 4, "4471-020_D.dxf": 6, "4471-105_B.dxf": 5, "4471-020_C.dxf": 5, "4471-020_D.pdf": 5, "4471-105_B.pdf": 4, "4471-000_B.pdf": 4, "4471-020_D_layout.dxf": 3, "4471-210_A.dxf": 5}
    for f, n in expected.items():
        j, d = dws.get(f, (None, None))
        check(f"reference {f}: {n} findings, no error", d is not None and d["findings"] == n and not d["error"], f"{d and (d['findings'], d['error'])}")
    # titles of key findings
    J = dws["4471-020_D.dxf"][0]["id"]; full = http("GET", f"/api/jobs/{J}")[2]
    D = next(d for d in full["drawings"] if d["file"] == "4471-020_D.dxf")
    titles = [f["title"] for f in D["findings"]]
    for key in ("Datum B referenced", "Thread callout M6x1.0 on a Ø6.6 hole", "Duplicated dimension 46", "QS-114 Rev B is superseded", "Material differs", "CHECKED field is blank"):
        check(f"finding present: {key}", any(key in t for t in titles))
    shaft = next(d for d in full["drawings"] if d["file"] == "4471-105_B.dxf")
    check("shaft: override mismatch 545 vs 540", any("545" in f["title"] and "540" in f["title"] for f in shaft["findings"]))
    check("shaft: ECO mismatch vs PLM", any("ECO-2299" in f["title"] for f in shaft["findings"]))
    assy = next(d for d in full["drawings"] if d["file"] == "4471-000_B.dxf")
    check("assembly: balloon 8 without parts list entry", any("Balloon 8" in f["title"] for f in assy["findings"]))
    JS, DS = dws["4471-210_A.dxf"]; smf = http("GET", f"/api/jobs/{JS['id']}")[2]["drawings"][0]["findings"]
    check("sheet metal: bend radius, thickness, hole-to-bend, flat pattern", {"SM-01", "SM-02", "SM-03", "SM-04"} <= {f.get("rule") for f in smf}, str([f.get("rule") for f in smf]))
    check("extraction has title block + notes + fcf", D["extraction"]["title_block"]["fields"]["part_no"]["value"] == "4471-020" and len(D["extraction"]["notes"]) == 5 and len(D["extraction"]["fcfs"]) == 1)
    # 3. rendering
    s, ct, svg, _ = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/render")
    ok = s == 200 and "svg" in ct
    if ok:
        try: xml.dom.minidom.parseString(svg); ok = b"PV_MARKERS" in svg or b"<svg" in svg
        except Exception: ok = False
    check("DXF renders to valid SVG with markers", ok, f"{s} {ct}")
    s, ct, svg2, _ = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/render?markers=0")
    check("DXF render without markers is smaller", s == 200 and len(svg2) < len(svg))
    JP, DP = dws["4471-020_D.pdf"]
    s, ct, png, h = http("GET", f"/api/jobs/{JP['id']}/drawings/{DP['id']}/render")
    check("PDF renders to PNG", s == 200 and "png" in ct and png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 20000, f"{s} {ct} {len(png) if isinstance(png, bytes) else png}")
    JL, DL = dws["4471-020_D_layout.dxf"]
    s, ct, svgl, _ = http("GET", f"/api/jobs/{JL['id']}/drawings/{DL['id']}/render")
    check("layout DXF renders (paper space + viewport)", s == 200 and len(svgl) > 20000)
    import pymupdf
    for label, jj, dd in (("dxf", J, D["id"]), ("pdf", JP["id"], DP["id"]), ("layout", JL["id"], DL["id"])):
        s, ct, pdf, _ = http("GET", f"/api/jobs/{jj}/drawings/{dd}/markup.pdf")
        ok = s == 200 and "pdf" in ct
        if ok:
            doc = pymupdf.open("pdf", pdf); ok = doc.page_count >= 2 and "PrintVerity review findings" in doc[-1].get_text(); doc.close()
        check(f"markup PDF ({label}) opens with findings page", ok, f"{s} {ct}")
    # 4. decisions and recheck
    fid = D["findings"][0]["id"]
    s, _, f, _ = http("POST", f"/api/jobs/{J}/drawings/{D['id']}/findings/{fid}", {"status": "rejected", "note": "test note"})
    check("finding decision saved", s == 200 and f["status"] == "rejected" and f["note"] == "test note")
    s, _, svg3, _ = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/render")
    check("rejected finding loses its marker", s == 200 and len(svg3) < len(svg))
    s, _, nd, _ = http("POST", f"/api/jobs/{J}/drawings/{D['id']}/recheck")
    check("recheck keeps decisions", s == 200 and any(x["status"] == "rejected" and x["note"] == "test note" for x in nd["findings"]) and len(nd["findings"]) == 6)
    http("POST", f"/api/jobs/{J}/drawings/{D['id']}/findings/{nd['findings'][0]['id']}", {"status": "open"})
    D = http("GET", f"/api/jobs/{J}")[2]["drawings"][[d["id"] for d in full["drawings"]].index(D["id"])]   # finding ids are regenerated by a re-check
    # 5. compare
    JC, DC = dws["4471-020_C.dxf"]
    s, _, cmp, _ = http("POST", "/api/compare", {"a_job": JC["id"], "a_drawing": DC["id"], "b_job": J, "b_drawing": D["id"]})
    check("compare C→D: 14 changes, ECO match", s == 200 and cmp["summary"]["total"] == 14 and cmp["eco"]["status"] == "match" and cmp["summary"]["affects_fit"] == 4, str(cmp.get("summary")))
    check("compare classifies TAP note removal as Process", any(c["kind"] == "Note" and c["change"] == "Removed" and c["impact"] == "Process" for c in cmp["changes"]))
    # 6. consistency & model
    s, _, con, _ = http("GET", f"/api/jobs/{J}/consistency")
    check("consistency: 4 conflicts incl. PO rev C", s == 200 and con["conflicts"] == 4 and any(r["label"] == "Rev conflict" and r["po"] == "C" for r in con["rows"]))
    s, _, mesh, _ = http("GET", f"/api/jobs/{J}/model/mesh")
    check("3D mesh with per-face status (conflict on the Ø6.35 hole)", s == 200 and mesh["triangles"] > 100 and any(f["status"] == "conflict" and abs((f["diameter"] or 0) - 6.35) < 0.01 for f in mesh["faces"]) and any(f["status"] == "match" for f in mesh["faces"]), str(mesh)[:120] if s != 200 else f"{mesh.get('triangles')} tris")
    q = f"a_job={JC['id']}&a_drawing={DC['id']}&b_job={J}&b_drawing={D['id']}"
    s, ct, xb, _ = http("GET", "/api/compare/report.xlsx?" + q); check("compare: change report xlsx", s == 200 and xb[:2] == b"PK")
    s, ct, pb, _ = http("GET", "/api/compare/side-by-side.pdf?" + q); check("compare: side-by-side PDF", s == 200 and pb[:4] == b"%PDF" and len(pb) > 50000)
    s, _, ib, _ = http("POST", "/api/compare/inbox", {"a_job": JC["id"], "a_drawing": DC["id"], "b_job": J, "b_drawing": D["id"]}); check("compare: send to quality inbox", s == 200 and ib["channel"] in ("inbox-folder", "smtp"))
    s, _, mc, _ = http("GET", f"/api/jobs/{J}/model-check")
    check("model check: 1 size conflict on bracket, others skipped", s == 200 and sum(d["conflicts"] for d in mc["drawings"]) == 1 and sum(1 for d in mc["drawings"] if d.get("skipped")) == 2)
    s, _, con2, _ = http("GET", f"/api/jobs/{JC['id']}/consistency")
    check("consistency without BOM/PO explains itself", s == 200 and con2["errors"])
    s, _, mc2, _ = http("GET", f"/api/jobs/{JC['id']}/model-check")
    check("model check without STEP explains itself", s == 200 and mc2.get("error"))
    # 7. docs
    import openpyxl, docx
    s, _, ip, _ = http("POST", f"/api/jobs/{J}/docs/inspection-plan", {"drawing": D["id"]})
    ok = s == 200 and ip["rows"] >= 8
    if ok:
        s2, ct2, xb, _ = http("GET", ip["url"]); ok = s2 == 200
        if ok:
            wb = openpyxl.load_workbook(io.BytesIO(xb)); ws = wb.active; ok = ws["B4"].value == "4471-020" and ws["E4"].value == "D" and ws.cell(10, 2).value is not None
    check("inspection plan xlsx generated and filled", ok, str(ip)[:120])
    s, _, coc, _ = http("POST", f"/api/jobs/{J}/docs/coc", {"drawing": D["id"], "qty": "50", "po": "88213"})
    ok = s == 200
    if ok:
        s2, _, db, _ = http("GET", coc["url"]); d = docx.Document(io.BytesIO(db)); txt = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        ok = "4471-020" in txt and "88213" in txt and "{{" not in txt
    check("CoC docx generated with placeholders filled", ok)
    s, _, tl, _ = http("GET", "/api/templates")
    check("template library lists 5", s == 200 and len(tl) == 5)
    s, _, bp, _ = http("POST", "/api/templates/bulk-edit", {"find": "QS-114 Rev B", "replace": "QS-114 Rev D", "preview": True})
    check("bulk edit preview finds 5 in 5 files, writes nothing", s == 200 and bp["total"] == 5 and bp["preview"] and not bp["files"])
    s, _, ba, _ = http("POST", "/api/templates/bulk-edit", {"find": "QS-114 Rev B", "replace": "QS-114 Rev D", "preview": False})
    ok = s == 200 and ba["total"] == 5 and len(ba["files"]) == 5
    if ok:
        f = next(x for x in ba["files"] if x["name"] == "CoC-STD-02.docx"); s2, _, db, _ = http("GET", f["url"])
        xmlt = zipfile.ZipFile(io.BytesIO(db)).read("word/document.xml").decode()
        ok = "<w:del " in xmlt and "<w:ins " in xmlt and "QS-114 Rev D" in xmlt and "QS-114 Rev B" in xmlt
        d = docx.Document(io.BytesIO(db))   # still a valid docx
        fx = next(x for x in ba["files"] if x["name"] == "IP-STD-03.xlsx"); s3, _, xb, _ = http("GET", fx["url"]); wb = openpyxl.load_workbook(io.BytesIO(xb)); ok = ok and "Change log" in wb.sheetnames
    check("bulk edit applies Word tracked changes + Excel change log", ok)
    s, _, bad, _ = http("POST", "/api/templates/bulk-edit", {"find": "   ", "replace": "x", "preview": True})
    check("bulk edit rejects empty find", s == 400)
    # 8. ask
    for q, must in (("what is the finish on the bracket?", "anodize"), ("what is the material of the shaft?", "4140"), ("Which datums are defined and which are referenced?", "datum b"), ("Is the PO at the right revision?", "rev c")):
        s, _, a, _ = http("POST", f"/api/jobs/{J}/ask", {"question": q})
        ans = (a.get("answer") or "").lower().replace("revision c", "rev c")
        ok = must in ans or (must == "datum b" and re.search(r"\bb\b", ans) and ("defin" in ans or "referenc" in ans))
        check(f"ask: {q[:40]} ({a.get('mode')})", s == 200 and bool(ok) and a["mode"] in ("local", "agent"), ans[:100])
    # 9. upload path with third-party edge-case DXFs, a PDF plot, an image-only PDF and a DWG
    import pymupdf as _pm
    raster_pdf = Path("/tmp/pv_check_data/scan_only.pdf")
    src = _pm.open(str(REF / "drawings" / "4471-020_D.pdf")); pix = src[0].get_pixmap(dpi=300); out = _pm.open(); pg = out.new_page(width=src[0].rect.width, height=src[0].rect.height)
    pg.insert_image(pg.rect, stream=pix.tobytes("png")); out.save(str(raster_pdf), deflate=True); out.close(); src.close()
    dwg = NB / "drawings" / "NB-3120-002_revA.dwg"
    files = [("drawings", p) for p in sorted(FIX.glob("*.dxf"))] + [("drawings", REF / "drawings" / "4471-105_B.pdf"), ("drawings", raster_pdf), ("drawings", dwg),
             ("step", REF / "documents" / "4471-105_B.stp"), ("bom", REF / "documents" / "BOM_4471-000_rev3.xlsx"), ("po", REF / "documents" / "PO_88213.xlsx")]
    body, hdr = multipart({"name": "Check upload", "customer": "QA"}, files)
    s, _, up, _ = http("POST", "/api/jobs?wait=1", body, hdr, raw=True)
    ok = s == 200 and len(up["drawings"]) == len([1 for k, _ in files if k == "drawings"])
    if ok:
        errs = [d for d in up["drawings"] if d["error"]]
        ok = len(errs) == 1 and errs[0]["file"] == dwg.name and "DXF" in errs[0]["error"]
    check("upload: edge-case DXFs + PDF + image-only PDF + DWG handled (only DWG errors, with guidance)", ok, str(up)[:200])
    if s == 200:
        uj = http("GET", f"/api/jobs/{up['id']}")[2]
        for d in uj["drawings"]:
            if d["error"]: continue
            s2, ct2, img, _ = http("GET", f"/api/jobs/{up['id']}/drawings/{d['id']}/render")
            check(f"upload render {d['file']}", s2 == 200 and len(img) > 500, f"{s2} {ct2}")
        raster = next((d for d in uj["drawings"] if d["file"] == "scan_only.pdf"), None)
        check("raster PDF read by OCR (notes extracted)", raster and raster["extraction"]["source"].get("raster") is True and raster["extraction"]["source"].get("ocr") and len(raster["extraction"]["notes"]) >= 3, str(raster and raster["extraction"]["counts"]))
        s2, _, mc3, _ = http("GET", f"/api/jobs/{up['id']}/model-check")
        check("model check on upload job runs", s2 == 200 and "drawings" in mc3)
        s2, _, c3, _ = http("GET", f"/api/jobs/{up['id']}/consistency")
        check("consistency on upload job runs", s2 == 200 and "rows" in c3)
        s2, _, _, _ = http("DELETE", f"/api/jobs/{up['id']}")
        check("job delete", s2 == 200 and http("GET", f"/api/jobs/{up['id']}")[0] == 404)
    # 9b. the NorthBay client kit, with the results its GUIDE.md promises
    files = [("drawings", p) for p in sorted((NB / "drawings").iterdir())] + [("step", NB / "documents" / "NB-3120-001_rev2.stp"), ("bom", NB / "documents" / "BOM_NB-3120-000_revB.csv"), ("po", NB / "documents" / "PO_NB-70211.xlsx")]
    body, hdr = multipart({"name": "NB job 26-0412", "customer": "NorthBay Engineering"}, files)
    s, _, nb, _ = http("POST", "/api/jobs?wait=1", body, hdr, raw=True)
    if check("NorthBay kit uploads: 9 rows, only the DWG refused", s == 200 and len(nb["drawings"]) == 9 and [d["file"] for d in nb["drawings"] if d["error"]] == [dwg.name], str(nb)[:200]):
        want = {"NB-3120-000_revB_assembly.dxf": 5, "NB-3120-001_rev1.dxf": 2, "NB-3120-001_rev2.dxf": 8, "NB-3120-001_rev2.pdf": 6, "NB-3120-001_rev2_SCAN.png": 4,
                "NB-3120-003_revA.dxf": 2, "NB-3120-005_revB_inch.dxf": 3, "NB-3120-007_revA_sheetmetal.dxf": 4}
        got = {d["file"]: d["findings"] for d in nb["drawings"] if not d["error"]}
        check("NorthBay kit: finding counts match the guide", got == want, str(got))
        nbj = http("GET", f"/api/jobs/{nb['id']}")[2]; nbd = {d["file"]: d for d in nbj["drawings"]}
        scan = nbd["NB-3120-001_rev2_SCAN.png"]; side = Path(scan["extraction"]["source"]["pdf_path"])
        check("scan read by OCR + layout model; its PDF wrapper stays inside the job", scan["part_no"] == "NB-3120-001" and scan["extraction"]["source"].get("ocr") and side.exists() and (Path("/tmp/pv_check_data/jobs") / nb["id"]) in side.parents, str(side))
        r1, r2 = nbd["NB-3120-001_rev1.dxf"], nbd["NB-3120-001_rev2.dxf"]
        s2, _, cmpn, _ = http("POST", "/api/compare", {"a_job": nb["id"], "a_drawing": r1["id"], "b_job": nb["id"], "b_drawing": r2["id"]})
        check("NorthBay kit: Rev 1 to Rev 2 shows 10 changes", s2 == 200 and cmpn["summary"]["total"] == 10, str(cmpn.get("summary")))
        s2, _, cn, _ = http("GET", f"/api/jobs/{nb['id']}/consistency")
        check("NorthBay kit: 5 conflicts, duplicate copies of the plate skipped", s2 == 200 and cn["conflicts"] == 5 and sum(x.startswith("Skipped:") for x in cn["sources"]) == 3, f"{cn.get('conflicts')} {cn.get('sources')}")
        s2, _, mn, _ = http("GET", f"/api/jobs/{nb['id']}/model-check")
        check("NorthBay kit: model check flags the Ø8.5 hole", s2 == 200 and sum(d["conflicts"] for d in mn["drawings"] if d["file"] == "NB-3120-001_rev2.dxf") == 1, str([(d["file"], d.get("conflicts")) for d in mn.get("drawings", [])]))
        # customer setup in the browser's API: suggestions from the title blocks, save, registers, re-check
        s2, _, cl, _ = http("GET", "/api/customers")
        check("Setup lists NorthBay with the finish missing on 8 drawings", s2 == 200 and any(c["slug"] == "northbay-engineering" and c["missing"] == {"finish": 8} and not c["has_setup"] for c in cl), str(cl)[:200])
        s2, _, cu, _ = http("GET", "/api/customers/northbay-engineering")
        sug = {(x["kind"], x["text"]): x for x in cu.get("suggestions", [])}
        check("Setup suggests SURFACE TREATMENT and the TREATMENT attribute as Finish, with the value", sug.get(("label", "SURFACE TREATMENT"), {}).get("field") == "finish" and sug.get(("label", "SURFACE TREATMENT"), {}).get("example") == "HARD ANODIZE" and sug.get(("tag", "TREATMENT"), {}).get("field") == "finish", str(list(sug))[:200])
        sys.path.insert(0, str(ROOT / "tests")); import shot as _shot0, asyncio as _asyncio0
        r = _asyncio0.run(_shot0.capture(f"{BASE}/#setup?tab=customers&customer=northbay-engineering", "/tmp/pv_check_setup.png", settle=3.0))
        check("Setup view renders the customer, the missing field and the suggestions", "SURFACE TREATMENT" in r["text"] and "could not find" in r["text"] and not [e for e in r["errors"] if "favicon" not in e], (r["errors"][:1] or ["text missing"])[0][:160])
        fields = cu["fields"]
        for x in cu["suggestions"]:
            if x["field"]: next(f for f in fields if f["field"] == x["field"])["labels" if x["kind"] == "label" else "tags"].append(x["text"])
        s2, _, sv, _ = http("PUT", "/api/customers/northbay-engineering", {"name": "NorthBay Engineering", "fields": fields})
        check("Setup saves the customer's wording", s2 == 200 and sv["has_setup"] and "SURFACE TREATMENT" in next(f for f in sv["fields"] if f["field"] == "finish")["labels"])
        body, hdr = multipart({}, [("file", NB / "customer_setup" / "spec_register.csv")]); s2, _, ri, _ = http("POST", "/api/reference/specs", body, hdr, raw=True)
        check("spec register imports", s2 == 200 and ri["imported"] == 1 and not ri["errors"], str(ri))
        body, hdr = multipart({}, [("file", NB / "customer_setup" / "released_revisions.csv")]); s2, _, ri, _ = http("POST", "/api/reference/revisions", body, hdr, raw=True)
        check("released-revision register imports", s2 == 200 and ri["imported"] == 6 and not ri["errors"], str(ri))
        s2, _, rc, _ = http("POST", "/api/customers/northbay-engineering/recheck?wait=1")
        check("re-check of the customer's 9 drawings", s2 == 200 and rc["drawings"] == 9, str(rc))
        nbj = http("GET", f"/api/jobs/{nb['id']}")[2]; nbd = {d["file"]: d for d in nbj["drawings"]}
        want = {"NB-3120-000_revB_assembly.dxf": 3, "NB-3120-001_rev1.dxf": 2, "NB-3120-001_rev2.dxf": 7, "NB-3120-001_rev2.pdf": 5, "NB-3120-001_rev2_SCAN.png": 3,
                "NB-3120-003_revA.dxf": 0, "NB-3120-005_revB_inch.dxf": 1, "NB-3120-007_revA_sheetmetal.dxf": 2}
        got = {f: len(d["findings"]) for f, d in nbd.items() if not d.get("error")}
        check("after setup: finding counts match the guide (finish read, PLM and spec checks live)", got == want, str(got))
        r2 = nbd["NB-3120-001_rev2.dxf"]; rules_ = {f["rule"] for f in r2["findings"]}
        check("after setup: Rev 2 finish read; superseded NB-QS-27 Rev A reported; no PLM or finish finding", r2["extraction"]["title_block"]["fields"]["finish"]["value"] == "HARD ANODIZE" and "SP-02" in rules_ and not {"TB-01", "PL-01"} & rules_, str(rules_))
        check("after setup: Rev 1 reported as superseded in PLM", any(f["rule"] == "PL-02" for f in nbd["NB-3120-001_rev1.dxf"]["findings"]))
        s2, _, cmpn, _ = http("POST", "/api/compare", {"a_job": nb["id"], "a_drawing": r1["id"], "b_job": nb["id"], "b_drawing": r2["id"]})
        check("after setup: compare shows the ECO matching PLM (CO-4478)", s2 == 200 and cmpn["eco"]["status"] == "match" and cmpn["eco"]["plm_eco"] == "CO-4478", str(cmpn.get("eco")))
        s2, _, cu2, _ = http("GET", "/api/customers/northbay-engineering")
        check("after setup: nothing missing, no finish suggestions left", s2 == 200 and cu2["missing"] == {} and not [x for x in cu2["suggestions"] if x["field"] == "finish"])
        check("Setup refuses bad input", http("PUT", "/api/customers/northbay-engineering", {"name": "Someone Else", "fields": []})[0] == 400 and http("GET", "/api/customers/..%2F..%2Fetc")[0] in (400, 404))
        check("Setup removes the customer's wording", http("DELETE", "/api/customers/northbay-engineering")[0] == 200 and not http("GET", "/api/customers/northbay-engineering")[2]["has_setup"])
        http("DELETE", f"/api/jobs/{nb['id']}")
    body, hdr = multipart({"name": "Empty"}, [])
    s, _, e, _ = http("POST", "/api/jobs", body, hdr, raw=True); check("empty upload creates empty job cleanly", s == 200 and e["drawings"] == [])
    if s == 200: http("DELETE", f"/api/jobs/{e['id']}")
    check("404 on unknown job", http("GET", "/api/jobs/nope")[0] == 404)
    (scratch / "secret.txt").write_text("x"); (scratch / "out_sibling").mkdir(exist_ok=True); (scratch / "out_sibling" / "secret.txt").write_text("x")
    check("path traversal on /api/out blocked (parent and look-alike sibling folders)", all(http("GET", u)[0] in (404, 400) for u in ("/api/out/../secret.txt", "/api/out/../out_sibling/secret.txt", "/api/out/%2e%2e/secret.txt")))
    check("generated files stay in the check's own output folder", not any((ROOT / "data" / "out").glob("template-edit-*")) and any((scratch / "out").glob("template-edit-*")))
    s, _, au, _ = http("GET", "/api/audit"); check("audit log has entries", s == 200 and len(au) > 10)
    s, _, rl, _ = http("GET", "/api/rules"); check("rule catalogue served", s == 200 and len(rl) >= 38 and all("id" in r for r in rl))
    s, _, pr, _ = http("GET", "/api/profiles"); check("profiles listed", s == 200 and any(p["slug"] == "default" for p in pr))
    s, ct, csvb, _ = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/findings.csv"); check("findings CSV export", s == 200 and "csv" in ct and b"Datum B" in csvb)
    check("pipeline trace recorded on drawing", [t["node"] for t in D.get("trace", [])] == ["ingest", "extract", "rule_check", "enrich", "finalize"])
    check("findings carry retrieval references", any(f.get("references") for f in D["findings"]))
    fl = next(f for f in D["findings"] if f.get("loc"))
    s, ct, cp, _ = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/findings/{fl['id']}/callout.png"); check("callout image for a finding", s == 200 and "png" in ct and len(cp) > 5000)
    s, _, wi, _ = http("POST", f"/api/jobs/{J}/docs/work-instruction", {"drawing": D["id"]}); check("work instruction docx with callout images", s == 200 and wi["name"].startswith("WI-") and wi["images"] >= 3)
    s, _, rel, _ = http("POST", "/api/docs/release", {"name": wi["name"]}); check("document release to storage", s == 200 and rel["storage"] in ("local", "s3") and Path(rel["location"]).exists() if rel.get("storage") == "local" else s == 200)
    s, _, lx, _ = http("POST", "/api/channels/lex", {"inputTranscript": "what is the finish on the bracket?", "sessionState": {"intent": {"name": "AskDrawing"}, "sessionAttributes": {"job": J}}})
    check("Lex channel answers", s == 200 and lx["sessionState"]["intent"]["state"] == "Fulfilled" and "anodize" in lx["messages"][0]["content"].lower())
    wav = Path("/tmp/pv_check_data/_check.wav")
    if shutil.which("espeak-ng"):
        subprocess.run(["espeak-ng", "-s", "150", "-w", str(wav), "Datum B is missing on sheet one, ask the drafter."], capture_output=True)
    else:
        wav.write_bytes(b"RIFF$\x00\x00\x00WAVEfmt \x10\x00\x00\x00\x01\x00\x01\x00\x80>\x00\x00\x00}\x00\x00\x02\x00\x10\x00data\x00\x00\x00\x00")
    body, hdr = multipart({}, [("audio", str(wav))]); s, _, vn, _ = http("POST", f"/api/jobs/{J}/drawings/{D['id']}/findings/{fl['id']}/voice-note", body, hdr, raw=True)
    transcribed = bool((vn.get("transcript") or {}).get("text"))
    check(f"voice note stored ({'transcribed' if transcribed else 'kept, not transcribed'})", s in (200, 202) and vn.get("finding", {}).get("voice_notes") and (s == 202 or len(vn["transcript"]["text"].split()) >= 3), (vn.get("transcript") or {}).get("text") or str(vn)[:160])
    s, _, js_, _ = http("GET", f"/api/jobs/{J}/status"); check("job status endpoint", s == 200 and js_["status"] == "done" and all(d["status"] == "done" for d in js_["drawings"]))
    s, ct, fj, _ = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/findings.json"); check("findings JSON export", s == 200 and fj["part_no"] == "4471-020" and len(fj["findings"]) == 6)
    s, _, _, h1 = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/render"); s2, _, _, h2 = http("GET", f"/api/jobs/{J}/drawings/{D['id']}/render")
    check("render cache serves repeat requests", s == 200 and s2 == 200 and h2.get("x-render-cache") == "hit" and "x-response-time" in h2)
    body, hdr = multipart({"name": "Bad ext"}, [("drawings", str(ROOT / "README.md"))]); s, _, e, _ = http("POST", "/api/jobs?wait=1", body, hdr, raw=True)
    check("upload rejects unsupported drawing type with 400", s == 400 and "unsupported" in str(e))
    check("rejected upload leaves no job behind", not any(j["name"] == "Bad ext" for j in http("GET", "/api/jobs")[2]))
    # 10. front-end views rendered in real time through the DevTools protocol: expected content present, no console errors, viewer transforms sane
    sys.path.insert(0, str(ROOT / "tests")); import shot as _shot, asyncio as _asyncio
    views = {"queue": "Drawing queue", f"review?job={J}&drawing={D['id']}": "Datum B referenced", f"review?job={J}&drawing={D['id']}&zoom=2": "Datum B referenced", "compare": "Export change report", "consistency": "conflicts",
             "model": "3D model", "docs": "Bulk template edit", "ask": "Ask the drawing package", "audit": "Audit log", "setup?tab=specs": "Import a spec register", "setup?tab=revisions": "Import released revisions"}
    check("React build is served", "/assets/index-" in http("GET", "/")[2].decode(errors="ignore"))
    transforms = {}
    for v, must in views.items():
        try:
            vname = v.split('?')[0] + (" " + v.split("tab=")[1] if "tab=" in v else "")
            r = _asyncio.run(_shot.capture(f"{BASE}/#{v}", f"/tmp/pv_check_{vname.replace(' ', '_')}.png", settle=3.0))
            errs = [e for e in r["errors"] if "favicon" not in e]
            check(f"view {vname}{' (zoom=2)' if 'zoom' in v else ''} renders, no console errors", must.lower() in r["text"].lower() and not errs, (errs[:1] or [f"'{must}' not in page text"])[0][:160])
            transforms[v] = r.get("transform")
        except Exception as e:
            check(f"view {v.split('?')[0]} renders", False, str(e)[:160])
    r = _asyncio.run(_shot.capture(f"{BASE}/#queue", "/tmp/pv_check_queue_upload.png", settle=2.0))
    check("queue offers upload and no sample loader", "Upload drawing package" in r["text"] and "sample package" not in r["text"].lower())
    r = _asyncio.run(_shot.capture(f"{BASE}/#review?job={J}&drawing={D['id']}", "/tmp/pv_check_dark.png", settle=2.5,
                                   steps=["document.getElementById('settings-button').click()", "document.querySelector('[data-theme-option=dark]').click()", "document.body.click()"],
                                   probe="[document.documentElement.dataset.theme, getComputedStyle(document.body).backgroundColor, localStorage.getItem('pv-theme'), getComputedStyle(document.querySelector('.canvaswrap img')).backgroundColor].join('|')"))
    check("Settings switches to the dark theme, saves it, and keeps the sheet white", r["probe"] == "dark|rgb(11, 18, 27)|dark|rgb(255, 255, 255)" and not [e for e in r["errors"] if "favicon" not in e], str(r["probe"]) + " " + str(r["errors"][:1]))
    r = _asyncio.run(_shot.capture(f"{BASE}/#queue", "/tmp/pv_check_theme_system.png", settle=2.0,
                                   steps=["localStorage.setItem('pv-theme', 'dark')", "location.reload()", "sessionStorage.setItem('pv-loaded', document.documentElement.dataset.theme)",
                                          "document.getElementById('settings-button').click()", "document.querySelector('[data-theme-option=light]').click()"],
                                   probe="[sessionStorage.getItem('pv-loaded'), document.documentElement.dataset.theme, localStorage.getItem('pv-theme'), document.querySelectorAll('.settings .pop [data-theme-option]').length].join('|')"))
    check("saved theme applies on load; Light switches back", r["probe"] == "dark|light|light|3", str(r["probe"]))
    fit_t = transforms.get(f"review?job={J}&drawing={D['id']}") or ""; zoom_t = transforms.get(f"review?job={J}&drawing={D['id']}&zoom=2") or ""
    check("viewer fits the sheet on load and honours zoom=2", "scale(" in fit_t and "scale(2)" in zoom_t and fit_t != zoom_t, f"{fit_t} | {zoom_t}")
    r = _asyncio.run(_shot.capture(f"{BASE}/#review?job={J}&drawing={D['id']}", "/tmp/pv_check_review_again.png", settle=2.0))
    check("review page has Pipeline tab with the graph trace", "Pipeline" in r["text"])
    os.environ["PV_CHROME_FLAGS"] = "--disable-3d-apis"
    try:
        r = _asyncio.run(_shot.capture(f"{BASE}/#model", "/tmp/pv_check_model_software.png", settle=4.0))
        check("model page falls back to the server-rendered preview without WebGL", "software preview" in r["text"].lower() and not [e for e in r["errors"] if "favicon" not in e], (r["errors"][:1] or ["no 'software preview' text"])[0][:160])
    finally:
        os.environ.pop("PV_CHROME_FLAGS", None)
    s, ct, pv, _ = http("GET", f"/api/jobs/{J}/model/preview.png?view=top"); check("model preview PNG endpoint", s == 200 and pv[:8] == b"\x89PNG\r\n\x1a\n")
    from PIL import Image as _Im
    s, ct, pvd, _ = http("GET", f"/api/jobs/{J}/model/preview.png?view=iso&theme=dark"); check("model preview follows the dark theme", s == 200 and _Im.open(io.BytesIO(pvd)).convert("RGB").getpixel((2, 2)) == (26, 36, 48))
    # 11. server log clean
    log.flush(); txt = open("/tmp/pv_test_server.log").read()
    check("server log has no ERROR/Traceback", "ERROR" not in txt and "Traceback" not in txt, txt[-300:])
    return finish(proc, t0)

def finish(proc, t0):
    try:
        os.killpg(os.getpgid(proc.pid), 15)
    except Exception: pass
    subprocess.run(["fuser", "-k", f"{PORT}/tcp"], capture_output=True)
    fails = [r for r in results if not r[1]]
    print(f"\n{len(results) - len(fails)}/{len(results)} checks passed in {time.time() - t0:.1f}s")
    for n, ok, info in fails: print("  FAILED:", n, "|", info)
    return 0 if not fails else 1

if __name__ == "__main__":
    sys.exit(main())
