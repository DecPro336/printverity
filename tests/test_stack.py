"""Tests for the orchestration, retrieval, geometry, vision, queue and integration layers. Run: .venv/bin/python -m pytest -q tests/"""
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import pytest
REF = ROOT / "testdata" / "reference"; DW = REF / "drawings"; DOC = REF / "documents"   # the reference test package

# ---------------- orchestration
def test_review_graph_trace_and_findings():
    from app.orchestration import run_review
    r = run_review(str(DW / "4471-020_D.dxf"))
    assert r["part_no"] == "4471-020" and len(r["findings"]) == 6
    assert [t["node"] for t in r["trace"]] == ["ingest", "extract", "rule_check", "enrich", "finalize"] and all(t["status"] == "ok" for t in r["trace"])
    assert any(f.get("references") for f in r["findings"])

def test_review_graph_declines_dwg(tmp_path):
    from app.orchestration import run_review
    p = tmp_path / "x.dwg"; p.write_bytes(b"AC1032" + b"\x00" * 64)
    r = run_review(str(p)); assert r["error"] and "DXF" in r["error"] and r["trace"][0]["status"] == "error"

def test_agent_local_planner():
    from app.orchestration.agent import PrintVerityAgent
    from app.orchestration import run_review
    job = {"id": "t", "name": "t", "drawings": [run_review(str(DW / "4471-020_D.dxf"))]}
    ag = PrintVerityAgent(); ag.client = None
    a = ag.ask(job, "what is the finish on the bracket?"); assert "ANODIZE" in a["answer"] and a["mode"] == "local"
    b = ag.ask(job, "what does Y14.5 say about datum references?"); assert "datum" in b["answer"].lower() and "search_knowledge" in b["tools_used"]
    c = ag.ask(job, "which datums are referenced?"); assert "Datum B" in c["answer"]

def test_agent_tool_schemas_match_functions():
    from app.orchestration.agent import TOOL_SCHEMAS, _tools
    assert set(t["name"] for t in TOOL_SCHEMAS) == set(_tools({"drawings": []}).keys())

# ---------------- retrieval
def test_hash_embedder_is_deterministic_and_normalized():
    from app.retrieval.embeddings import HashEmbedder
    e = HashEmbedder(64); v1, v2 = e.embed(["datum B missing", "datum B missing"])
    assert v1 == v2 and abs(sum(x * x for x in v1) - 1) < 1e-6

def test_local_store_roundtrip(tmp_path):
    from app.retrieval.store import LocalStore, Document
    from app.retrieval.embeddings import HashEmbedder
    e = HashEmbedder(64); st = LocalStore(tmp_path / "v.json", 64)
    docs = [Document("a", "spec", "QS-114", "anodize spec"), Document("b", "standard", "Y14.5 datums", "datum reference rules")]
    for d in docs: d.vector = e.embed([d.text])[0]
    assert st.upsert(docs) == 2 and st.count() == 2
    hit = st.search(e.embed(["datum rules"])[0], k=1)[0]; assert hit.id == "b"
    assert st.search(e.embed(["x"])[0], k=5, kinds=["spec"])[0].kind == "spec"

def test_index_documents_cover_sources():
    from app.retrieval import index as idx
    docs = idx.all_documents([]); kinds = {d.kind for d in docs}
    assert {"spec", "rule", "standard", "template"} <= kinds and any(d.id == "spec:QS-114" for d in docs)

def test_retriever_service_search():
    from app.retrieval import get_retriever
    r = get_retriever(); r.refresh([])
    hits = r.search("position tolerance references a datum that is not defined", k=2, kinds=["standard"])
    assert hits and "datum" in hits[0].text.lower()

# ---------------- geometry
def test_step_engine_reads_sample():
    from app.geometry import parse_step, engine_name
    st = parse_step(str(DOC / "4471-020_D.stp"))
    assert st["extents"]["x"] == 120.0 and any(abs(c["diameter"] - 6.6) < 0.01 and c["surfaces"] == 3 for c in st["cylinders"])
    assert st["engine"] in ("opencascade", "text-parser") and engine_name() in ("opencascade", "text-parser")

def test_step_text_parser_units_inch(tmp_path):
    from app.geometry import step_text as step_parse
    p = tmp_path / "in.stp"; p.write_text("ISO-10303-21;\nHEADER;\nFILE_SCHEMA(('AP214'));\nENDSEC;\nDATA;\n#1 = ( LENGTH_UNIT() NAMED_UNIT(*) CONVERSION_BASED_UNIT('INCH',#2) );\n#3 = CYLINDRICAL_SURFACE('',#4,0.125);\nENDSEC;\nEND-ISO-10303-21;\n")
    st = step_parse.parse(str(p)); assert st["units"] == "in" and abs(st["cylinders"][0]["diameter"] - 6.35) < 0.01

# ---------------- vision
def test_ocr_normalization():
    from app.vision.raster import normalize_ocr
    assert normalize_ocr("120 +/-0.1") == "120 ±0.1" and normalize_ocr("1 0F 1") == "1 OF 1" and normalize_ocr("O6.60 THRU").startswith("Ø")

def test_raster_pdf_reads_notes_and_title(tmp_path):
    import pymupdf
    from app.vision.raster import extract_raster_pdf
    src = pymupdf.open(str(DW / "4471-020_D.pdf")); pix = src[0].get_pixmap(dpi=300)
    out = pymupdf.open(); page = out.new_page(width=src[0].rect.width, height=src[0].rect.height); page.insert_image(page.rect, pixmap=pix); p = tmp_path / "scan.pdf"; out.save(str(p))
    ex = extract_raster_pdf(str(p))
    assert ex["source"]["raster"] and ex["source"]["ocr"] and len(ex["notes"]) >= 4
    assert ex["title_block"]["fields"].get("part_no", {}).get("value", "").replace(" ", "").startswith("4471")

def test_layout_model_interface():
    from app.vision.layout import get_layout_model, CLASSES
    m = get_layout_model(); assert m.name in ("yolov8", "heuristic") and "title_block" in CLASSES

# ---------------- queue and integrations
def test_worker_status_shape():
    from app.tasks import worker_status
    st = worker_status(); assert set(st) == {"mode", "broker", "workers", "async"}

def test_plm_service_local_fallback():
    from app.integrations.epicor import PlmService
    s = PlmService(); rec, src = s.part_revision("4471-020"); assert rec["rev"] == "D" and src in ("local", "epicor")
    assert s.part_revision("nope")[0] is None

def test_lex_handler_roundtrip():
    from app.integrations import lex
    ev = {"inputTranscript": "what is the finish?", "sessionState": {"intent": {"name": "AskDrawing"}, "sessionAttributes": {"job": "j1"}}}
    r = lex.handle_event(ev, lambda job, q: {"answer": f"{job}:{q}", "citations": ["c1"]})
    assert r["sessionState"]["intent"]["state"] == "Fulfilled" and "j1:what is the finish?" in r["messages"][0]["content"]
    assert lex.handle_event({"inputTranscript": "x", "sessionState": {}}, lambda j, q: {})["sessionState"]["intent"]["state"] == "Failed"

def test_voice_notes_not_configured(tmp_path, monkeypatch):
    from app.integrations import voice_notes
    from app import config
    monkeypatch.setattr(config, "ASSEMBLYAI_API_KEY", "")
    with pytest.raises(voice_notes.NotConfigured): voice_notes.transcribe(str(tmp_path / "a.wav"))

def test_voice_notes_missing_file_with_key(tmp_path, monkeypatch):
    from app.integrations import voice_notes
    from app import config
    monkeypatch.setattr(config, "ASSEMBLYAI_API_KEY", "test-key")
    with pytest.raises(Exception): voice_notes.transcribe(str(tmp_path / "missing.wav"))

def test_storage_local(tmp_path):
    from app.integrations.storage import LocalStorage
    src = tmp_path / "doc.txt"; src.write_text("x"); st = LocalStorage(tmp_path / "rel")
    assert Path(st.put(str(src), "doc.txt")).read_text() == "x"

def test_callout_image(tmp_path):
    from app.orchestration import run_review
    from app.integrations.illustrations import callout_png
    d = run_review(str(DW / "4471-020_D.dxf")); f = next(x for x in d["findings"] if x.get("loc"))
    png = callout_png(d, f); assert png[:8] == b"\x89PNG\r\n\x1a\n" and len(png) > 5000

def test_work_instruction_doc():
    from app.orchestration import run_review
    from app.docs import fill_work_instruction
    d = run_review(str(DW / "4471-020_D.dxf")); r = fill_work_instruction(d, {"name": "T"}, {})
    assert Path(r["path"]).exists() and r["name"].startswith("WI-4471-020")

# ---------------- sheet metal, mesh, exports
def test_sheet_metal_rules():
    from app import dxf_extract, rules
    ex = dxf_extract.extract(str(DW / "4471-210_A.dxf")); ids = {f["rule"] for f in rules.run_all(ex)}
    assert {"SM-01", "SM-02", "SM-03", "SM-04"} <= ids
    assert "TH-03" not in ids            # CLS-M4-1 is a PEM part number, not a thread callout

def test_sheet_metal_rules_stay_quiet_on_machined_parts():
    from app import dxf_extract, rules
    ex = dxf_extract.extract(str(DW / "4471-020_D.dxf")); assert not any(f["rule"].startswith("SM-") for f in rules.run_all(ex))

def test_mesh_tessellation():
    from app.geometry.mesh import tessellate
    m = tessellate(str(DOC / "4471-020_D.stp"))
    assert m["triangles"] > 100 and any(f["kind"] == "cylinder" and abs((f["diameter"] or 0) - 6.6) < 0.01 for f in m["faces"]) and m["bbox"][1][0] == 120.0

def test_compare_exports(tmp_path, monkeypatch):
    from app.orchestration import run_review
    from app.compare import compare
    from app import exports
    monkeypatch.setattr(exports, "INBOX", tmp_path / "inbox")
    a = run_review(str(DW / "4471-020_C.dxf")); b = run_review(str(DW / "4471-020_D.dxf")); rep = compare(a["extraction"], b["extraction"], "Rev C", "Rev D")
    x = exports.change_report_xlsx(rep); assert x.exists() and x.suffix == ".xlsx"
    pdf = exports.side_by_side_pdf(a, b, rep); assert pdf[:4] == b"%PDF"
    r = exports.send_to_quality_inbox(rep, [x]); assert r["channel"] == "inbox-folder" and (tmp_path / "inbox").exists()

def test_inch_thread_on_major_diameter_hole():
    """Unified thread tolerances are physical (0.0025 in), so a 5/16-18 callout on a Ø0.3125 hole is caught while 1/4-20 on Ø0.201 passes."""
    from app import rules
    ex = {"source": {"kind": "dxf", "units": "in"}, "texts": [], "title_block": {"fields": {}},
          "threads": [{"spec": "5/16-18 UNC", "system": "unified", "major": 0.3125, "pitch": 1/18, "tap_drill": 0.257, "units": "in", "loc": [0, 0], "circle": {"r": 0.15625, "via": "leader"}},
                      {"spec": "1/4-20 UNC", "system": "unified", "major": 0.25, "pitch": 0.05, "tap_drill": 0.201, "units": "in", "loc": [0, 0], "circle": {"r": 0.1005, "via": "leader"}}]}
    ids = [f["rule"] for f in rules.run_all(ex) if f["category"] == "Thread / hole"]
    assert ids == ["TH-01"]

@pytest.mark.parametrize("name", ["circle_radius_le_0.dxf", "clockwise_arcs_hatch.dxf", "dimension_in_block.dxf"])
def test_edge_case_dxf_reviews_without_error(name):
    """Third-party DXFs from the ezdxf examples (degenerate circle, clockwise arcs with hatch, dimension inside a block)."""
    from app.orchestration import run_review
    d = run_review(str(ROOT / "tests" / "fixtures" / name))
    assert not d.get("error") and d["extraction"]["source"]["kind"] == "dxf" and isinstance(d["findings"], list)

def test_model_preview_light_and_dark():
    import io
    from PIL import Image
    from app.geometry.mesh import tessellate
    from app.geometry.preview import render, PALETTE
    m = tessellate(str(DOC / "4471-020_D.stp")); faces = [{**f, "status": "neutral"} for f in m["faces"]]
    for dark in (False, True):
        img = Image.open(io.BytesIO(render(m, faces, view="iso", dark=dark))).convert("RGB")
        assert img.getpixel((2, 2)) == PALETTE[dark][0]

def test_coc_blank_qty_and_po_keep_the_write_in_line():
    import docx
    from app.docs import fill_coc
    from app.orchestration import run_review
    d = run_review(str(DW / "4471-020_D.dxf"))
    r = fill_coc(d, {"name": "T"}, {"qty": "", "po": "  "})
    doc = docx.Document(r["path"]); txt = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert "____" in txt and "{{" not in txt

def test_search_waits_for_the_startup_warmup():
    import threading, time
    from app.retrieval import service
    r = service.get_retriever(); service.WARMED.clear(); threading.Timer(0.3, service.WARMED.set).start()
    t0 = time.time(); hits = r.search("datum referenced but not defined", k=1, kinds=["standard"])
    assert time.time() - t0 >= 0.25 and hits and service.WARMED.is_set()
