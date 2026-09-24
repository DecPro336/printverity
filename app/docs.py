"""Documentation: fill the client's Word/Excel templates from drawing data, and bulk-edit templates with tracked changes."""
import copy, datetime
from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Side
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from .util import DATA, OUT

TEMPLATES = DATA / "templates"

def _tb(ex, f):
    return ((ex.get("title_block", {}).get("fields", {}).get(f) or {}).get("value") or "").strip()

def _method(dim):
    t = (dim.get("text") or "").upper(); tol = dim.get("tol"); v = dim.get("value") or 0
    if dim.get("type") == "diameter" or "Ø" in t:
        if tol and (tol["plus"] + tol["minus"]) <= 0.05: return "Pin gage / bore gage"
        return "Pin gage" if v < 30 else "Bore gage"
    if dim.get("type") == "radius": return "Radius gage"
    if tol and (tol["plus"] + tol["minus"]) <= 0.1: return "CMM"
    if tol and (tol["plus"] + tol["minus"]) <= 0.3: return "Micrometer"
    return "Caliper"

def _freq(dim):
    tol = dim.get("tol")
    if tol and (tol["plus"] + tol["minus"]) <= 0.1: return "100%"
    if tol: return "1 per 10"
    return "First + 1 per 25"

def characteristics(ex):
    rows = []; n = 1
    for d in ex.get("dimensions", []):
        if d.get("value") is None: continue
        tol = d.get("tol")
        tol_s = f"±{tol['plus']:g}" if tol and tol["plus"] == tol["minus"] else (f"+{tol['plus']:g}/-{tol['minus']:g}" if tol else "General (note)")
        name = {"diameter": "Diameter", "radius": "Radius", "angular": "Angle", "linear": "Length", "aligned": "Length"}.get(d.get("type"), "Dimension")
        if "SLOT" in (d.get("text") or "").upper(): name = "Slot"
        rows.append({"n": n, "characteristic": f"{name} {d.get('text')}", "nominal": f"{d['value']:g}" if d.get("type") != "angular" else d.get("text"), "tol": tol_s, "method": _method(d), "freq": _freq(d), "source": "Sheet 1, view"})
        n += 1
    for f in ex.get("fcfs", []):
        rows.append({"n": n, "characteristic": f"{(f.get('symbol') or 'GD&T').capitalize()} to {'-'.join(f.get('datums', [])) or '-'}", "nominal": ("Ø" if f.get("diameter") else "") + f"{f.get('tolerance')}", "tol": " ".join(f.get("modifiers", [])) or "—", "method": "CMM", "freq": "100%", "source": "Sheet 1, FCF"})
        n += 1
    for th in ex.get("threads", []):
        rows.append({"n": n, "characteristic": f"Thread {th['spec']}", "nominal": th["spec"], "tol": th.get("fit") or "6H/6g", "method": "Thread plug gage GO/NOGO", "freq": "100%", "source": "Sheet 1, callout"}); n += 1
    for nt in ex.get("notes", []):
        U = nt["text"].upper()
        if any(w in U for w in ("FINISH", "ANODIZE", "PLATE", "HEAT TREAT", "HARDEN", "PAINT", "POWDER", "BLACK OXIDE")):
            rows.append({"n": n, "characteristic": f"Note {nt['n']}: {nt['text'][:60]}", "nominal": "Per note", "tol": "—", "method": "Cert review", "freq": "Per lot", "source": f"Note {nt['n']}"}); n += 1
    return rows

def fill_inspection_plan(drawing, job):
    ex = drawing["extraction"]; pn = _tb(ex, "part_no") or Path(drawing["file"]).stem; rev = _tb(ex, "rev") or "x"
    src = TEMPLATES / "IP-STD-03.xlsx"
    wb = openpyxl.load_workbook(src); ws = wb.active
    ph = {"{{PART_NO}}": pn, "{{REV}}": rev, "{{TITLE}}": _tb(ex, "title"), "{{MATERIAL}}": _tb(ex, "material"), "{{JOB}}": job.get("name", ""), "{{DATE}}": datetime.date.today().isoformat(), "{{DRAWING_FILE}}": drawing["file"]}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "{{" in c.value:
                for k, v in ph.items(): c.value = c.value.replace(k, v)
    # find the header row of the characteristics table
    start = None
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 1).value).strip() == "#": start = r + 1; break
    if start is None: start = ws.max_row + 2
    rows = characteristics(ex)
    thin = Side(style="thin", color="BBBBBB"); border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, ch in enumerate(rows):
        r = start + i
        vals = [ch["n"], ch["characteristic"], ch["nominal"], ch["tol"], ch["method"], ch["freq"], ch["source"], ""]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(r, j, v); cell.border = border; cell.alignment = Alignment(vertical="top", wrap_text=True)
    out = OUT / f"IP-{pn}-{rev}.xlsx"
    wb.save(out)
    return {"path": str(out), "name": out.name, "rows": len(rows), "characteristics": rows}

def _replace_in_paragraph(p, mapping):
    text = "".join(r.text for r in p.runs)
    if not any(k in text for k in mapping): return
    for k, v in mapping.items(): text = text.replace(k, v)
    for r in p.runs[1:]: r.text = ""
    if p.runs: p.runs[0].text = text

def fill_coc(drawing, job, extra=None):
    ex = drawing["extraction"]; pn = _tb(ex, "part_no") or Path(drawing["file"]).stem; rev = _tb(ex, "rev") or "x"
    src = TEMPLATES / "CoC-STD-02.docx"
    d = docx.Document(src)
    specs = ", ".join(sorted({s["ref"] + (f" Rev {s['rev']}" if s.get("rev") else "") for s in ex.get("specs", [])})) or "per drawing"
    mapping = {"{{PART_NO}}": pn, "{{REV}}": rev, "{{TITLE}}": _tb(ex, "title"), "{{MATERIAL}}": _tb(ex, "material") or "per drawing", "{{FINISH}}": _tb(ex, "finish") or "per drawing",
               "{{JOB}}": job.get("name", ""), "{{DATE}}": datetime.date.today().isoformat(), "{{SPECS}}": specs, "{{QTY}}": str((extra or {}).get("qty") or "").strip() or "____", "{{PO}}": str((extra or {}).get("po") or "").strip() or "____",
               "{{ECO}}": _tb(ex, "eco") or "-"}
    for p in d.paragraphs: _replace_in_paragraph(p, mapping)
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs: _replace_in_paragraph(p, mapping)
    out = OUT / f"CoC-{pn}-{rev}.docx"
    d.save(out)
    return {"path": str(out), "name": out.name}

# ---------------------------------------------------------------- tracked-change bulk edit
def _tracked_replace_run(p, run, old, new, author, date, next_id):
    """Split a run around `old` and emit <w:del>old</w:del><w:ins>new</w:ins> in place, preserving run formatting."""
    text = run.text
    idx = text.find(old)
    if idx < 0: return False
    before, after = text[:idx], text[idx + len(old):]
    rpr = run._r.find(qn("w:rPr"))
    def mk_run(t, deleted=False):
        r = OxmlElement("w:r")
        if rpr is not None: r.append(copy.deepcopy(rpr))
        te = OxmlElement("w:delText" if deleted else "w:t"); te.set(qn("xml:space"), "preserve"); te.text = t; r.append(te)
        return r
    parent = run._r.getparent(); pos = parent.index(run._r)
    new_nodes = []
    if before: new_nodes.append(mk_run(before))
    de = OxmlElement("w:del"); de.set(qn("w:id"), str(next_id())); de.set(qn("w:author"), author); de.set(qn("w:date"), date); de.append(mk_run(old, True)); new_nodes.append(de)
    ie = OxmlElement("w:ins"); ie.set(qn("w:id"), str(next_id())); ie.set(qn("w:author"), author); ie.set(qn("w:date"), date); ie.append(mk_run(new)); new_nodes.append(ie)
    if after: new_nodes.append(mk_run(after))
    parent.remove(run._r)
    for i, nn in enumerate(new_nodes): parent.insert(pos + i, nn)
    return True

def _tracked_paragraph(p, old, new, author, date, next_id):
    count = 0
    # merge runs if the match spans run boundaries
    full = "".join(r.text for r in p.runs)
    if old not in full: return 0
    if not any(old in r.text for r in p.runs) and p.runs:
        for r in p.runs[1:]: r.text = ""
        p.runs[0].text = full
    guard = 0
    while guard < 50:
        guard += 1
        hit = False
        for r in list(p.runs):
            if old in r.text:
                if _tracked_replace_run(p, r, old, new, author, date, next_id): count += 1; hit = True; break
        if not hit: break
    return count

def _cell_replace(ws, old, new):
    n = 0; changed = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and old in c.value:
                c.value = c.value.replace(old, new); n += 1; changed.append(c.coordinate)
    return n, changed

def list_templates():
    out = []
    for p in sorted(TEMPLATES.glob("*")):
        if p.suffix.lower() in (".docx", ".xlsx"):
            out.append({"name": p.name, "kind": "Word" if p.suffix.lower() == ".docx" else "Excel", "size": p.stat().st_size, "description": TEMPLATE_INFO.get(p.name, "")})
    return out

TEMPLATE_INFO = {"IP-STD-03.xlsx": "Inspection plan", "CoC-STD-02.docx": "Certificate of conformance", "CP-STD-01.docx": "Control plan", "RTR-STD-04.xlsx": "Router / traveler", "PPAP-PSW.docx": "PPAP part submission warrant"}

def bulk_edit(find, replace, author="PrintVerity", preview_only=False):
    """Apply a find/replace across every template. Word files get tracked changes; Excel files are edited with a change log sheet."""
    date = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    results = []; files = []
    counter = [900]
    def next_id():
        counter[0] += 1; return counter[0]
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S")
    outdir = OUT / f"template-edit-{stamp}"
    if not preview_only: outdir.mkdir(parents=True, exist_ok=True)
    for tpl in sorted(TEMPLATES.glob("*")):
        if tpl.suffix.lower() == ".docx":
            d = docx.Document(tpl); n = 0; previews = []
            paras = list(d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    for cell in row.cells: paras.extend(cell.paragraphs)
            for p in paras:
                txt = "".join(r.text for r in p.runs)
                if find in txt:
                    previews.append({"before": txt.strip()[:160], "after": txt.replace(find, replace).strip()[:160]})
                    if not preview_only: n += _tracked_paragraph(p, find, replace, author, date, next_id)
                    else: n += txt.count(find)
            if n:
                if not preview_only:
                    outp = outdir / tpl.name; d.save(outp); files.append({"name": tpl.name, "path": str(outp), "kind": "Word", "mode": "tracked changes"})
                results.append({"file": tpl.name, "kind": "Word", "count": n, "previews": previews[:4], "mode": "tracked changes"})
        elif tpl.suffix.lower() == ".xlsx":
            wb = openpyxl.load_workbook(tpl); total = 0; previews = []; changed_all = []
            for ws in wb.worksheets:
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and find in c.value:
                            previews.append({"before": c.value[:160], "after": c.value.replace(find, replace)[:160], "cell": f"{ws.title}!{c.coordinate}"})
                if not preview_only:
                    n, changed = _cell_replace(ws, find, replace); total += n; changed_all += [f"{ws.title}!{x}" for x in changed]
                else:
                    total += len([p for p in previews if p["cell"].startswith(ws.title + "!")])
            if total:
                if not preview_only:
                    log = wb.create_sheet("Change log")
                    log.append(["Date", "Author", "Cell", "Find", "Replace"])
                    for cc in changed_all: log.append([date, author, cc, find, replace])
                    outp = outdir / tpl.name; wb.save(outp); files.append({"name": tpl.name, "path": str(outp), "kind": "Excel", "mode": "change log sheet"})
                results.append({"file": tpl.name, "kind": "Excel", "count": total, "previews": previews[:4], "mode": "change log sheet"})
    return {"find": find, "replace": replace, "results": results, "files": files, "total": sum(r["count"] for r in results), "outdir": str(outdir) if not preview_only else None, "preview": preview_only}


def fill_work_instruction(drawing, job, images):
    """Work instruction (docx): one section per open finding with a callout image and the corrective instruction.
    `images` maps finding id -> PNG bytes (see integrations.illustrations)."""
    from docx.shared import Mm
    ex = drawing["extraction"]; pn = _tb(ex, "part_no") or Path(drawing["file"]).stem; rev = _tb(ex, "rev") or "x"
    d = docx.Document()
    d.add_heading(f"Work instruction: {pn} Rev {rev}", 1)
    p = d.add_paragraph(); p.add_run(f"{_tb(ex, 'title')}   ·   Job {job.get('name', '')}   ·   Generated {datetime.date.today().isoformat()} from {drawing['file']}").italic = True
    d.add_paragraph("Each item below was raised by the drawing review. Resolve it with the drafter or the customer before the router is released.")
    for f in drawing.get("findings", []):
        if f.get("status") == "rejected": continue
        d.add_heading(f"{f.get('n', '')}. {f['title']}", 2)
        d.add_paragraph(f"Severity: {f['severity']}   Rule: {f.get('rule', '')}   Confidence: {f['confidence']:.2f}   Status: {f['status']}")
        d.add_paragraph(f["detail"])
        if f.get("references"):
            d.add_paragraph("Basis: " + "; ".join(f"{r['title']}" for r in f["references"][:2]))
        png = images.get(f["id"])
        if png:
            tmp = OUT / f"_wi_{f['id']}.png"; tmp.write_bytes(png)
            d.add_picture(str(tmp), width=Mm(110)); tmp.unlink(missing_ok=True)
        if f.get("note"): d.add_paragraph(f"Engineer note: {f['note']}")
    out = OUT / f"WI-{pn}-{rev}.docx"; d.save(out)
    return {"path": str(out), "name": out.name}
