"""Exports of the revision comparison: xlsx change report, side-by-side PDF, and the quality-inbox hand-off."""
from __future__ import annotations
import datetime, json, smtplib, logging, os
from email.message import EmailMessage
from pathlib import Path
import openpyxl, pymupdf
from openpyxl.styles import Font, PatternFill, Alignment
from . import render
from .util import OUT, now

log = logging.getLogger("printverity.exports")
INBOX = OUT / "inbox"

def change_report_xlsx(report: dict) -> Path:
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Change report"
    ws["A1"] = f"Revision change report: {report['labelA']} → {report['labelB']}"; ws["A1"].font = Font(bold=True, size=13)
    e = report.get("eco", {})
    ws["A2"] = f"Part {e.get('part_no')} · Drawing ECO {e.get('drawing_eco') or '—'} · PLM ECO {e.get('plm_eco') or '—'} ({e.get('status')}) · generated {now()}"
    hdr = ["#", "Kind", "Item", report["labelA"], report["labelB"], "Change", "Impact", "Why"]
    ws.append([]); ws.append(hdr)
    for c in ws[4]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="13294B")
    fills = {"Fit": "FBE3DF", "Process": "FDE9D3", "Inspection": "FDE9D3", "Purchasing": "FFF4CC", "Cosmetic": "E3EEFB"}
    for i, c in enumerate(report.get("changes", []), 1):
        ws.append([i, c["kind"], c["item"], c["a"], c["b"], c["change"], c["impact"], c.get("why", "")])
        ws.cell(ws.max_row, 7).fill = PatternFill("solid", fgColor=fills.get(c["impact"], "FFFFFF"))
    ws.append([]); ws.append(["Downstream actions"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    for d in report.get("downstream", []): ws.append(["", d])
    for col, w in zip("ABCDEFGH", (5, 12, 22, 30, 30, 14, 12, 50)): ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
    out = OUT / f"CHG-{(e.get('part_no') or 'part')}-{report['labelA'].split()[-1]}-to-{report['labelB'].split()[-1]}.xlsx"
    wb.save(out); return out

def _sheet_pdf(drawing: dict) -> pymupdf.Document:
    ex = drawing["extraction"]
    if drawing["kind"] == "dxf":
        return pymupdf.open("pdf", render.dxf_to_pdf_bytes(drawing["path"], [], ex["bbox"], ex.get("space", "model"), ex.get("viewports")))
    path = ex["source"].get("pdf_path") or drawing["path"]
    return pymupdf.open(path)

def side_by_side_pdf(a: dict, b: dict, report: dict) -> bytes:
    """Landscape pages: old revision left, new revision right, changed items listed under each pair, then the change table."""
    da, db = _sheet_pdf(a), _sheet_pdf(b)
    out = pymupdf.open()
    n = max(da.page_count, db.page_count)
    for i in range(n):
        page = out.new_page(width=1190, height=842)
        page.insert_text((30, 30), f"{report['labelA']}  (left)   vs   {report['labelB']}  (right)   ·   sheet {i + 1}", fontsize=11, fontname="hebo")
        for k, doc in enumerate((da, db)):
            if i < doc.page_count:
                rect = pymupdf.Rect(20 + k * 580, 45, 20 + k * 580 + 570, 45 + 570 * doc[i].rect.height / doc[i].rect.width)
                page.show_pdf_page(rect, doc, i)
        y = 640
        page.insert_text((30, y), "Changes affecting fit or process:", fontsize=10, fontname="hebo"); y += 14
        for c in [c for c in report.get("changes", []) if c["impact"] in ("Fit", "Process", "Inspection")][:10]:
            page.insert_text((40, y), f"{c['kind']}: {c['a']} → {c['b']}  [{c['impact']}] {c.get('why', '')}"[:150], fontsize=8.5, fontname="helv"); y += 11
    page = out.new_page(width=595, height=842); y = 50
    page.insert_text((40, y), "Change report", fontsize=16, fontname="hebo"); y += 24
    for i, c in enumerate(report.get("changes", []), 1):
        page.insert_text((40, y), f"{i}. [{c['impact']}] {c['kind']} · {c['item']}: {c['a']} → {c['b']} ({c['change']})"[:120], fontsize=9, fontname="helv"); y += 12
        if y > 800: page = out.new_page(width=595, height=842); y = 50
    data = out.tobytes(garbage=3, deflate=True); out.close(); da.close(); db.close()
    return data

def send_to_quality_inbox(report: dict, attachments: list[Path], sender: str = "PrintVerity") -> dict:
    """Hands the change report to the quality team: SMTP when PV_SMTP_HOST / PV_QUALITY_INBOX are set, otherwise a dated drop in out/inbox with a manifest."""
    INBOX.mkdir(parents=True, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d-%H%M%S"); e = report.get("eco", {})
    subject = f"[PrintVerity] Change report {e.get('part_no')} {report['labelA']} → {report['labelB']}: {report['summary']['total']} changes, {report['summary']['affects_fit']} affect fit"
    body = f"{subject}\n\nDownstream:\n" + "\n".join(f"- {d}" for d in report.get("downstream", [])) + "\n\nECO: drawing {} / PLM {} ({})\n".format(e.get("drawing_eco"), e.get("plm_eco"), e.get("status"))
    host, to = os.environ.get("PV_SMTP_HOST"), os.environ.get("PV_QUALITY_INBOX")
    if host and to:
        msg = EmailMessage(); msg["Subject"] = subject; msg["From"] = os.environ.get("PV_SMTP_FROM", "printverity@localhost"); msg["To"] = to; msg.set_content(body)
        for p in attachments:
            msg.add_attachment(Path(p).read_bytes(), maintype="application", subtype="octet-stream", filename=Path(p).name)
        with smtplib.SMTP(host, int(os.environ.get("PV_SMTP_PORT", "25")), timeout=10) as s:
            if os.environ.get("PV_SMTP_USER"): s.starttls(); s.login(os.environ["PV_SMTP_USER"], os.environ.get("PV_SMTP_PASSWORD", ""))
            s.send_message(msg)
        return {"channel": "smtp", "to": to, "subject": subject}
    drop = INBOX / f"{stamp}-{e.get('part_no') or 'part'}"; drop.mkdir(parents=True, exist_ok=True)
    (drop / "message.txt").write_text(body, encoding="utf-8")
    for p in attachments: (drop / Path(p).name).write_bytes(Path(p).read_bytes())
    (drop / "manifest.json").write_text(json.dumps({"subject": subject, "files": [Path(p).name for p in attachments], "time": now()}, indent=2), encoding="utf-8")
    return {"channel": "inbox-folder", "location": str(drop), "subject": subject, "note": "Set PV_SMTP_HOST and PV_QUALITY_INBOX to send by email"}
